"""Create local editable spreadsheets/slides and PDF reports, without Office automation."""
import io
import json
import math
import os
from pathlib import Path
import re
import uuid
from xml.sax.saxutils import escape

FOLDER = Path.home() / 'Documents' / 'JARVIS Documents'


def text(value, limit=10000):
    if not isinstance(value,str) or len(value)>limit:
        raise ValueError(f'Text must be at most {limit} characters.')
    return value


def table(spec):
    columns=spec.get('columns',[]); rows=spec.get('rows',[])
    if not isinstance(columns,list) or not 1<=len(columns)<=20 or any(not isinstance(c,str) or len(c)>100 for c in columns):
        raise ValueError('Supply 1–20 column names.')
    if not isinstance(rows,list) or len(rows)>2000:
        raise ValueError('Supply at most 2000 rows.')
    for row in rows:
        if not isinstance(row,list) or len(row)!=len(columns): raise ValueError('Every row must match the column count.')
        for cell in row:
            if cell is not None and not isinstance(cell,(str,int,float,bool)): raise ValueError('Cells must be text, numbers, booleans or null.')
            if isinstance(cell,str) and len(cell)>2000: raise ValueError('Cell text exceeds 2000 characters.')
            if isinstance(cell,(int,float)) and not math.isfinite(cell): raise ValueError('Numbers must be finite.')
    return columns,rows


def chart_data(spec):
    columns,rows=table(spec)
    if not rows or len(rows)>40: raise ValueError('Charts support 1–40 categories.')
    if len(columns)<2 or any(type(r[1]) not in (int,float) for r in rows):
        raise ValueError('Charts need category labels in column one and numeric values in column two.')
    return columns,[str(r[0]) for r in rows],[r[1] for r in rows]


def workbook(spec):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.utils import get_column_letter
    columns,rows=table(spec)
    wb=Workbook(); ws=wb.active; ws.title='Data'
    for row in [columns]+rows:
        ws.append(row)
    # Strings are literal: imported text can never become formulas, links or macros.
    for row in ws:
        for cell in row:
            if isinstance(cell.value,str): cell.data_type='s'
            cell.alignment=Alignment(vertical='top',wrap_text=True)
    for cell in ws[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='164E63')
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    totals={}
    for c,heading in enumerate(columns,1):
        letter=get_column_letter(c); ws.column_dimensions[letter].width=min(48,max(16,len(heading)+3))
        if spec.get('totals') and rows and all(type(row[c-1]) in (int,float) for row in rows):
            ws.cell(len(rows)+2,c,f'=SUM({letter}2:{letter}{len(rows)+1})')
            totals[heading]=sum(row[c-1] for row in rows)
    if spec.get('chart'):
        chart_data(spec)
        chart=LineChart() if spec['chart']=='line' else BarChart()
        chart.title=spec['title']; chart.add_data(Reference(ws,min_col=2,min_row=1,max_row=len(rows)+1),titles_from_data=True)
        chart.set_categories(Reference(ws,min_col=1,min_row=2,max_row=len(rows)+1))
        ws.add_chart(chart,f'A{len(rows)+5}')
    output=io.BytesIO(); wb.save(output); wb.close()
    return output.getvalue(), {'rows':len(rows),'totals':totals,'formulas':'Generated totals recalculate when opened in Excel or LibreOffice.'}


def pdf(spec):
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.graphics.shapes import Drawing
    from reportlab.graphics.charts.barcharts import VerticalBarChart
    from reportlab.graphics.charts.linecharts import HorizontalLineChart
    styles=getSampleStyleSheet()
    font=Path(os.environ.get('WINDIR','C:/Windows'))/'Fonts/arial.ttf'
    if font.is_file():
        if 'JarvisUnicode' not in pdfmetrics.getRegisteredFontNames(): pdfmetrics.registerFont(TTFont('JarvisUnicode',str(font)))
        for style in styles.byName.values(): style.fontName='JarvisUnicode'
    story=[Paragraph(escape(spec['title']),styles['Title']),Spacer(1,12)]
    content=text(spec.get('content',''),60000)
    for line in content.splitlines():
        heading=line.startswith('# ')
        story.append(Paragraph(escape(line[2:] if heading else line) or ' ',styles['Heading2' if heading else 'BodyText']))
        story.append(Spacer(1,5))
    if spec.get('columns'):
        columns,rows=table(spec)
        if len(columns)>8 or len(rows)>500: raise ValueError('PDF tables support at most 8 columns and 500 rows.')
        grid=[[Paragraph(escape(str(v)) if v is not None else '',styles['BodyText']) for v in row] for row in [columns]+rows]
        block=Table(grid,repeatRows=1,colWidths=[468/len(columns)]*len(columns))
        block.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#d7edf2')),
            ('GRID',(0,0),(-1,-1),.4,colors.lightgrey),('VALIGN',(0,0),(-1,-1),'TOP'),('BOTTOMPADDING',(0,0),(-1,-1),7)]))
        story.extend([Spacer(1,12),block])
    if spec.get('chart'):
        _,labels,values=chart_data(spec)
        chart=HorizontalLineChart() if spec['chart']=='line' else VerticalBarChart()
        chart.x=45; chart.y=60; chart.width=390; chart.height=170
        chart.data=[values]; chart.categoryAxis.categoryNames=labels
        chart.valueAxis.valueMin=min(0,min(values)); chart.valueAxis.valueMax=max(0,max(values)) or 1
        if spec['chart']=='line': chart.lines[0].strokeColor=colors.HexColor('#0891b2')
        else: chart.bars[0].fillColor=colors.HexColor('#0891b2')
        chart.categoryAxis.labels.angle=30; chart.categoryAxis.labels.fontSize=8
        drawing=Drawing(468,260); drawing.add(chart); story.append(drawing)
    def footer(canvas,doc):
        canvas.setFont('Helvetica',9); canvas.drawRightString(540,30,f'JARVIS · {doc.page}')
    output=io.BytesIO()
    SimpleDocTemplate(output,title=spec['title'],author='JARVIS',leftMargin=72,rightMargin=72).build(story,onFirstPage=footer,onLaterPages=footer)
    return output.getvalue(),{'kind':'report'}


def slides(spec):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    entries=spec.get('slides')
    if not isinstance(entries,list) or not 1<=len(entries)<=40: raise ValueError('Supply 1–40 slides.')
    def textbox(slide,value,x,y,w,h,size,color='183B4E'):
        shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); shape.text=text(value,6000)
        shape.text_frame.word_wrap=True
        for p in shape.text_frame.paragraphs: p.font.size=Pt(size); p.font.color.rgb=RGBColor.from_string(color)
    for index,entry in enumerate(entries):
        if not isinstance(entry,dict): raise ValueError('Each slide must be an object.')
        title=text(entry.get('title',''),140); body=text(entry.get('body',''),1800)
        # Split long prose into continuation slides to avoid silent clipping.
        import textwrap
        parts=textwrap.wrap(body,width=350 if entry.get('chart') else 650,replace_whitespace=False) or ['']
        for page,part in enumerate(parts):
            slide=prs.slides.add_slide(prs.slide_layouts[6])
            textbox(slide,title+(' (continued)' if page else ''),.65,.45,12,1.1,30)
            has_chart=bool(entry.get('chart')) and page==0
            textbox(slide,part,.7,1.7,5.2 if has_chart else 11.9,4.9,22)
            if has_chart:
                columns,labels,values=chart_data(entry)
                data=CategoryChartData(); data.categories=labels; data.add_series(columns[1],values)
                kind=XL_CHART_TYPE.LINE if entry['chart']=='line' else XL_CHART_TYPE.COLUMN_CLUSTERED
                slide.shapes.add_chart(kind,Inches(6.1),Inches(1.7),Inches(6.5),Inches(4.8),data)
            textbox(slide,f'{spec["title"]}  |  {len(prs.slides)}',.7,7,12,.3,10,'53727F')
    output=io.BytesIO(); prs.save(output)
    return output.getvalue(),{'slides':len(prs.slides)}


def create(fmt,spec):
    if fmt not in ('xlsx','pdf','pptx') or not isinstance(spec,dict): raise ValueError('Choose xlsx, pdf or pptx and supply an object.')
    if len(json.dumps(spec))>300000: raise ValueError('Artifact request is too large.')
    title=text(spec.get('title',''),120).strip()
    if not title: raise ValueError('A title is required.')
    if spec.get('chart') not in (None,False,'bar','line'): raise ValueError('Chart must be bar or line.')
    data,details={'xlsx':workbook,'pdf':pdf,'pptx':slides}[fmt](dict(spec,title=title))
    FOLDER.mkdir(parents=True,exist_ok=True)
    safe=re.sub(r'[^\w .-]','',title).strip(' .')[:65] or 'Document'
    name='Jarvis-'+safe+'-'+uuid.uuid4().hex[:10]+'.'+fmt
    path=FOLDER/name
    with path.open('xb') as handle: handle.write(data)
    return {'path':str(path),'name':name,'format':fmt,**details}


def catalog():
    if not FOLDER.exists(): return []
    return [{'name':p.name,'bytes':p.stat().st_size} for p in sorted(FOLDER.glob('Jarvis-*'),key=lambda p:p.stat().st_mtime,reverse=True)
            if p.is_file() and not p.is_symlink() and p.suffix in ('.xlsx','.pdf','.pptx')][:100]


def open_file(name):
    if name not in {p['name'] for p in catalog()}: raise ValueError('Choose a generated artifact.')
    path=FOLDER/name
    if not path.resolve().is_relative_to(FOLDER.resolve()): raise ValueError('Invalid artifact path.')
    os.startfile(str(path)); return {'opened':name}
