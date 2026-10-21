import datetime as dt
import json
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import patch, Mock
from zoneinfo import ZoneInfo

import artifacts
import digests
import memory
import privacy


class LocalFeaturesTest(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup);self.root=Path(self.temp.name)
        for module,field,value in [(privacy,'STORE',self.root/'privacy.json'),(memory,'PATH',str(self.root/'memory.json')),
                (memory,'RECENT_PATH',str(self.root/'recent.json')),(digests,'STORE',self.root/'digests.json'),
                (digests,'REPORTS',self.root/'reports'),(artifacts,'FOLDER',self.root/'files')]:
            p=patch.object(module,field,value);p.start();self.addCleanup(p.stop)

    def test_redacts_credentials_json_assignments_and_selected_information(self):
        privacy.configure(['Secret Project'],email=True,phone=True)
        for secret in ['sk-proj-'+'a'*30,'ghp_'+'B'*30,'123-45-6789','Bearer abcdefghijklmnop',
                       'eyJabc.abcdefgh.signature','user@example.com','312-555-1234','Secret Project']:
            self.assertNotIn(secret,privacy.redact('prefix '+secret+' suffix'))
        self.assertNotIn('hunter22',privacy.redact('{"password":"hunter22"}'))
        self.assertIn('player progress',privacy.redact('player progress'))

    def test_private_key_and_image_payload(self):
        key='-----BEGIN PRIVATE KEY-----\nabc123\n-----END PRIVATE KEY-----'
        self.assertNotIn('abc123',privacy.redact(key))
        self.assertEqual(privacy.payload({'type':'image','source':{'type':'base64','data':'sk-'+'a'*30}})['source']['data'],'sk-'+'a'*30)
        self.assertNotIn('sk-'+'a'*30,privacy.payload({'data':'sk-'+'a'*30})['data'])

    def test_memory_redaction_edit_and_stale_revision(self):
        memory.remember('password=hunter22')
        memory.record_turn('SSN 123-45-6789','okay')
        state=memory.snapshot(); self.assertNotIn('hunter22',state['facts'][0]['fact'])
        self.assertNotIn('123-45-6789',state['recent'][0]['user'])
        updated=memory.edit('update',state['revision'],state['facts'][0]['id'],'Likes tea')
        with self.assertRaises(ValueError):memory.edit('delete',state['revision'],updated['facts'][0]['id'])
        deleted=memory.edit('delete',updated['revision'],updated['facts'][0]['id'])
        self.assertEqual(deleted['facts'],[])
        cleared=memory.edit('clear_recent',deleted['revision']);self.assertEqual(cleared['recent'],[])

    def test_memory_scrub_rewrites_existing_records(self):
        memory.remember('private phrase');memory.record_turn('private phrase','private phrase')
        privacy.configure(['private phrase'])
        memory.edit('scrub',memory.revision())
        self.assertNotIn('private phrase',Path(memory.PATH).read_text())
        self.assertNotIn('private phrase',Path(memory.RECENT_PATH).read_text())
        self.assertEqual(memory.recent_prompt(),'')

    def test_spoken_forget_also_invalidates_old_context(self):
        memory.remember('Likes tea');memory.record_turn('Likes tea','Saved')
        memory.forget('Likes tea')
        self.assertEqual(memory.facts(),[]);self.assertEqual(memory.recent_prompt(),'')

    def test_memory_edit_refreshes_active_and_restart_context(self):
        import jarvis
        memory.remember('Likes tea');memory.record_turn('Likes tea','Tea noted')
        brain=jarvis.Brain();brain._run=lambda speak:speak('Okay')
        self.assertIn('Likes tea',brain.system)
        state=memory.snapshot();memory.edit('update',state['revision'],state['facts'][0]['id'],'Likes coffee')
        brain.ask('What do I prefer?',Mock())
        self.assertIn('Likes coffee',brain.system);self.assertNotIn('Likes tea',brain.system)
        self.assertNotIn('Likes tea',jarvis.Brain().system)

    def test_openai_boundary_redacts_without_touching_authentication(self):
        import model_provider
        response=Mock(status_code=200)
        secret='sk-proj-'+'Q'*30
        with patch.object(model_provider,'load_openai_key',return_value='real-auth'), \
                patch.object(model_provider.requests,'post',return_value=response) as post:
            model_provider.Responses()._post({'model':'test','input':secret})
        self.assertNotIn(secret,str(post.call_args.kwargs['json']))
        self.assertEqual(post.call_args.kwargs['headers']['Authorization'],'Bearer real-auth')

    def test_codex_connection_redacts_model_text_only(self):
        import agent_sessions
        from types import SimpleNamespace
        conn=agent_sessions.Connection.__new__(agent_sessions.Connection)
        conn.lock=threading.Lock();conn.proc=SimpleNamespace(stdin=Mock())
        conn.send({'method':'turn/start','params':{'threadId':'thread','input':[{'type':'text','text':'password=hunter22'}]}})
        sent=json.loads(conn.proc.stdin.write.call_args.args[0])
        self.assertNotIn('hunter22',str(sent));self.assertEqual(sent['params']['threadId'],'thread')

    def job(self,**kwargs):
        return digests.upsert(dict(name='Test',kind='progress',at='18:00',enabled=True,**kwargs))

    def test_disabled_by_default_and_quiet_hours(self):
        self.assertEqual(digests.load()['jobs'],[])
        job=self.job();worker=Mock()
        quiet=dt.datetime(2026,10,21,23,tzinfo=ZoneInfo('America/Chicago'))
        self.assertEqual(digests.run(job['id'],automatic=True,now=quiet,worker=worker)['status'],'not_due')
        worker.assert_not_called()
        early=quiet.replace(hour=17)
        self.assertEqual(digests.run(job['id'],automatic=True,now=early,worker=worker)['status'],'not_due')

    def test_restart_dedupe_daily_limit_and_report_sources(self):
        job=self.job();now=dt.datetime(2026,10,21,19,tzinfo=ZoneInfo('America/Chicago'))
        result={'text':'Evidence','sources':[{'title':'Source','url':'https://example.com'}],'model_calls':0,'searches':0}
        worker=Mock(return_value=result)
        done=digests.run(job['id'],automatic=True,now=now,worker=worker)
        self.assertEqual(digests.report(done['id'])['sources'],result['sources'])
        self.assertEqual(digests.run(job['id'],automatic=True,now=now,worker=worker)['status'],'already_attempted')
        digests.run(job['id'],now=now,worker=worker)
        with self.assertRaises(ValueError):digests.run(job['id'],now=now,worker=worker)
        self.assertEqual(worker.call_count,2)

    def test_failure_and_restart_are_not_automatic_retries(self):
        job=self.job();now=dt.datetime(2026,10,21,19,tzinfo=ZoneInfo('America/Chicago'))
        failed=digests.run(job['id'],now=now,worker=Mock(side_effect=RuntimeError('injected')))
        self.assertEqual(failed['status'],'failed')
        self.assertEqual(digests.run(job['id'],automatic=True,now=now)['status'],'already_attempted')
        data=digests.load();data['runs'][0]['status']='running';digests.save(data)
        stop=digests.start();stop.set()
        self.assertEqual(digests.load()['runs'][0]['status'],'interrupted')

    def test_polling_worker_dispatches_only_enabled_schedules(self):
        job=self.job()
        digests.upsert({'name':'disabled','kind':'research','at':'18:00','enabled':False})
        event=Mock();event.wait.side_effect=[False,True];event.is_set.return_value=False
        def thread(**kwargs):
            return Mock(start=lambda:kwargs['target']())
        with patch.object(digests.threading,'Event',return_value=event), \
                patch.object(digests.threading,'Thread',side_effect=thread), patch.object(digests,'run') as run:
            digests.start()
        run.assert_called_once_with(job['id'],automatic=True)

    def test_schedule_validation_and_report_path_boundary(self):
        with self.assertRaises(ValueError):digests.configure({'daily_limit':0})
        with self.assertRaises(ValueError):digests.configure({'timezone':'not-a-zone'})
        with self.assertRaises(ValueError):digests.upsert({'name':'bad','kind':'shell','at':'18:00'})
        with self.assertRaises(ValueError):digests.report('../outside')
        with self.assertRaises(ValueError):digests.clock('25:00')

    def test_digest_sources_exclude_uncited_search_results(self):
        cited={'url':'https://example.com/update','title':'Update'}
        other={'url':'https://example.org/unrelated','title':'Other'}
        self.assertEqual(digests.source_evidence('Source: https://example.com/update',[other,cited,cited]),[cited])

    def test_xlsx_roundtrip_totals_chart_and_formula_injection(self):
        from openpyxl import load_workbook
        spec={'title':'Sample budget','columns':['Category','Amount'],'rows':[['Art',40],['Audio',25]],'totals':True,'chart':'bar'}
        result=artifacts.create('xlsx',spec)
        self.assertEqual(result['totals']['Amount'],65)
        wb=load_workbook(result['path']);ws=wb.active
        self.assertEqual(ws['B4'].value,'=SUM(B2:B3)');self.assertEqual(len(ws._charts),1);wb.close()
        result=artifacts.create('xlsx',dict(spec,rows=[['=HYPERLINK("https://example.com","x")',3]],chart=False))
        wb=load_workbook(result['path']);self.assertEqual(wb.active['A2'].data_type,'s');wb.close()

    def test_pdf_roundtrip_multiline_table_and_literal_markup(self):
        from pypdf import PdfReader
        result=artifacts.create('pdf',{'title':'Sample report','content':'# Progress\nLiteral <script> & text\n'+'A long report paragraph.\n'*100,
            'columns':['Category','Amount'],'rows':[['Art',40],['Audio',25]],'chart':'bar'})
        reader=PdfReader(result['path']);self.assertGreater(len(reader.pages),1)
        text=''.join(p.extract_text() for p in reader.pages)
        self.assertIn('Literal <script> & text',text);self.assertIn('Audio',text)

    def test_slides_roundtrip_chart_and_continuations(self):
        from pptx import Presentation
        result=artifacts.create('pptx',{'title':'Sample slides','slides':[{'title':'Budget','body':'Progress update. '*75,
            'columns':['Category','Amount'],'rows':[['Art',40],['Audio',25]],'chart':'bar'}]})
        prs=Presentation(result['path']);self.assertGreater(len(prs.slides),1)
        self.assertEqual(sum(bool(s.has_chart) for slide in prs.slides for s in slide.shapes),1)
        for slide in prs.slides:
            for shape in slide.shapes:
                self.assertLessEqual(shape.left+shape.width,prs.slide_width)
                self.assertLessEqual(shape.top+shape.height,prs.slide_height)

    def test_artifacts_no_overwrite_or_arbitrary_open(self):
        spec={'title':'../../CON','columns':['A'],'rows':[['x']]}
        first=artifacts.create('xlsx',spec);second=artifacts.create('xlsx',spec)
        self.assertNotEqual(first['path'],second['path'])
        self.assertEqual(Path(first['path']).parent,artifacts.FOLDER)
        with self.assertRaises(ValueError):artifacts.open_file('../outside.pdf')
        with self.assertRaises(ValueError):artifacts.create('exe',spec)
        with self.assertRaises(ValueError):artifacts.create('xlsx',dict(spec,rows=[[float('nan')]]))


if __name__=='__main__':unittest.main()
